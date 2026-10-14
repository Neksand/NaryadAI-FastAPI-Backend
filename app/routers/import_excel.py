"""Excel-импорт справочников equipment и materials.

POST /api/v1/dict/{ctype}/import (multipart, только admin).

- dry_run=true  -> preview {"rows": [...], "errors": [...]} без записи.
- dry_run=false -> атомарное применение в одной транзакции (всё или ничего),
  ответ {"created": N, "updated": M}. Любая построчная ошибка блокирует запись.

Чистая функция парсинга :func:`parse_import_workbook` вынесена отдельно,
чтобы её можно было тестировать без живой БД.
"""

from __future__ import annotations

import io
import uuid
from typing import Any

from fastapi import APIRouter, Depends, File, Query, UploadFile
from openpyxl import load_workbook

from app.db import get_pool
from app.deps import get_current_user, require_role
from app.errors import not_found, validation_error
from app.routers.admin import CATALOGUES

router = APIRouter()

IMPORTABLE_TYPES = ("equipment", "materials")

REQUIRED_COLUMNS: dict[str, list[str]] = {
    # NOT NULL-колонки из migrations/001_initial.sql.
    "equipment": ["name", "inventory_no", "area_id", "type", "criticality"],
    "materials": ["name", "unit"],
}

# Природный уникальный ключ справочника: по нему ищем дубли внутри файла
# и различаем created/updated при применении.
UNIQUE_COLUMN: dict[str, str] = {
    "equipment": "inventory_no",
    "materials": "name",
}

EQUIPMENT_CRITICALITY = ("A", "B", "C")
MATERIAL_UNITS = ("шт", "кг", "л", "м")


def _norm(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value if value != "" else None
    return value


def _err(row: int, path: str, message: str) -> dict[str, Any]:
    return {"row": row, "path": path, "message": message}


def parse_import_workbook(
    data: bytes, ctype: str
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Разобрать xlsx-байты в (rows, errors).

    Номера строк — excel-номера (заголовок = 1, первая строка данных = 2).
    Каждая ошибка — dict с ключами row/path/message.
    Невалидный файл -> ValueError (HTTP-слой превращает в 422).
    """
    if ctype not in IMPORTABLE_TYPES:
        raise ValueError(f"Неизвестный справочник: {ctype}")
    if not data:
        raise ValueError("Пустой файл")
    try:
        wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError("Не удалось прочитать xlsx-файл") from e
    ws = wb.active
    if ws is None:
        raise ValueError("В книге нет листов")
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        raise ValueError("Файл пустой: нет даже заголовка")
    raw_header = all_rows[0]
    header: list[str] = [
        str(c).strip() for c in raw_header if c is not None and str(c).strip() != ""
    ]
    if not header:
        raise ValueError("Пустой заголовок")
    if len(set(header)) != len(header):
        seen: set[str] = set()
        dup = ""
        for h in header:
            if h in seen:
                dup = h
                break
            seen.add(h)
        raise ValueError(f"Дублирующаяся колонка в заголовке: {dup}")

    allowed = set(CATALOGUES[ctype]["cols"])
    errors: list[dict[str, Any]] = []
    for col in REQUIRED_COLUMNS[ctype]:
        if col not in header:
            errors.append(_err(1, col, f"Нет обязательной колонки: {col}"))
    if errors:
        return [], errors

    rows: list[dict[str, Any]] = []
    seen_keys: set[Any] = set()
    seen_qr: set[Any] = set()
    key_col = UNIQUE_COLUMN[ctype]
    # Позиции колонок по заголовку (включая неизвестные — их игнорируем ниже).
    positions = list(raw_header)

    for excel_row, values in enumerate(all_rows[1:], start=2):
        vals = list(values) + [None] * (len(positions) - len(values))
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in vals):
            continue  # пропускаем полностью пустые строки
        record: dict[str, Any] = {}
        for i, h in enumerate(positions):
            name = str(h).strip() if h is not None else ""
            if name and name in allowed:
                record[name] = _norm(vals[i]) if i < len(vals) else None
        row_errors: list[dict[str, Any]] = []
        for col in REQUIRED_COLUMNS[ctype]:
            if record.get(col) is None:
                row_errors.append(_err(excel_row, col, "Обязательное поле пустое"))
        key = record.get(key_col)
        if key is not None:
            if key in seen_keys:
                row_errors.append(_err(excel_row, key_col, f"Дублирующийся код в файле: {key}"))
            else:
                seen_keys.add(key)
        if ctype == "equipment":
            crit = record.get("criticality")
            if crit is not None and crit not in EQUIPMENT_CRITICALITY:
                row_errors.append(
                    _err(excel_row, "criticality", "criticality должен быть A, B или C")
                )
            area = record.get("area_id")
            if area is not None:
                try:
                    uuid.UUID(str(area))
                except ValueError:
                    row_errors.append(_err(excel_row, "area_id", "area_id должен быть UUID"))
            qr = record.get("qr_token")
            if qr is not None:
                if qr in seen_qr:
                    row_errors.append(
                        _err(excel_row, "qr_token", f"Дублирующийся qr_token в файле: {qr}")
                    )
                else:
                    seen_qr.add(qr)
        if ctype == "materials":
            unit = record.get("unit")
            if unit is not None and unit not in MATERIAL_UNITS:
                row_errors.append(_err(excel_row, "unit", "unit должен быть шт, кг, л или м"))
            usage = record.get("typical_usage_per_order")
            if usage is not None and not isinstance(usage, (int, float)):
                try:
                    record["typical_usage_per_order"] = float(str(usage).replace(",", "."))
                except ValueError:
                    row_errors.append(
                        _err(
                            excel_row,
                            "typical_usage_per_order",
                            "typical_usage_per_order должен быть числом",
                        )
                    )
        if row_errors:
            errors.extend(row_errors)
        else:
            rows.append(record)
    return rows, errors


@router.post("/dict/{ctype}/import")
async def import_dict(
    ctype: str,
    file: UploadFile = File(...),
    dry_run: bool = Query(default=True),
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    require_role(user, "admin")
    if ctype not in IMPORTABLE_TYPES:
        raise not_found("Справочник не найден")
    content = await file.read()
    try:
        rows, errors = parse_import_workbook(content, ctype)
    except ValueError as e:
        raise validation_error([{"path": "file", "message": str(e)}]) from e
    if dry_run:
        return {"rows": rows, "errors": errors}
    if errors:
        raise validation_error(errors)
    pool = await get_pool()
    meta = CATALOGUES[ctype]
    table = str(meta["table"])
    key_col = UNIQUE_COLUMN[ctype]
    cols = [c for c in meta["cols"] if rows and c in rows[0]]
    set_cols = [c for c in cols if c != key_col]
    created = 0
    updated = 0
    async with pool.acquire() as conn, conn.transaction():
        for record in rows:
            existing = await conn.fetchrow(
                f"SELECT id FROM {table} WHERE {key_col}=$1", record.get(key_col)
            )
            if existing is not None:
                if set_cols:
                    sets = ", ".join(f"{c} = ${i + 1}" for i, c in enumerate(set_cols))
                    extra = (
                        ", updated_at = now(), deleted_at = NULL"
                        if ctype == "equipment"
                        else ", deleted_at = NULL"
                    )
                    await conn.execute(
                        f"UPDATE {table} SET {sets}{extra} WHERE {key_col}=${len(set_cols) + 1}",
                        *[record[c] for c in set_cols],
                        record.get(key_col),
                    )
                updated += 1
            else:
                placeholders = ", ".join(f"${i + 1}" for i in range(len(cols)))
                await conn.execute(
                    f"INSERT INTO {table}({','.join(cols)}) VALUES ({placeholders})",
                    *[record[c] for c in cols],
                )
                created += 1
    return {"created": created, "updated": updated}
