"""Генерация XLSX/PDF, порт src/reports/report.service.ts."""

import io
import json
import os
import uuid
from typing import Any


async def _rows(conn: Any, export: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    rtype = export["report_type"]
    filters = (
        export["filters"]
        if isinstance(export["filters"], dict)
        else json.loads(export["filters"] or "{}")
    )
    scope = export.get("scope")
    if scope is None:
        scope = {}
    if isinstance(scope, str):
        scope = json.loads(scope or "{}")
    if not isinstance(scope, dict):
        raise ValueError("Report scope is invalid")
    role = scope.get("role")
    if role not in {"worker", "master", "manager", "admin"}:
        raise ValueError("Report scope has no valid role")

    # Revalidate persisted jobs as well as API input. This protects the worker from
    # malformed/legacy JSON and keeps all values bound as SQL parameters.
    from app.routers.reports import ExportFilters

    parsed = ExportFilters.model_validate(filters)
    from_at = parsed.from_
    to_at = parsed.to
    area_ids = None
    worker_id = None
    crew_id = None
    if role == "master":
        try:
            area_ids = [uuid.UUID(str(value)) for value in scope.get("area_ids", [])]
        except (TypeError, ValueError):
            area_ids = []  # A corrupt master scope fails closed.
    elif role == "worker":
        try:
            worker_id = uuid.UUID(str(scope["user_id"]))
            crew_id = uuid.UUID(str(scope["crew_id"])) if scope.get("crew_id") else None
        except (KeyError, TypeError, ValueError):
            worker_id = uuid.UUID(int=0)
            crew_id = None

    args = (
        area_ids,
        from_at,
        to_at,
        parsed.area_id,
        parsed.equipment_id,
        parsed.assignee_id,
        parsed.status,
        parsed.priority,
        parsed.kind,
        parsed.work_order_id,
        worker_id,
        crew_id,
    )
    where = """($1::uuid[] IS NULL OR w.area_id=ANY($1::uuid[]))
        AND ($2::timestamptz IS NULL OR w.issued_at >= $2::timestamptz)
        AND ($3::timestamptz IS NULL OR w.issued_at < $3::timestamptz)
        AND ($4::uuid IS NULL OR w.area_id=$4::uuid)
        AND ($5::uuid IS NULL OR w.equipment_id=$5::uuid)
        AND ($6::uuid IS NULL OR w.assignee_id=$6::uuid OR w.crew_id=$6::uuid)
        AND ($7::work_order_status IS NULL OR w.status=$7::work_order_status)
        AND ($8::work_order_priority IS NULL OR w.priority=$8::work_order_priority)
        AND ($9::work_order_kind IS NULL OR w.kind=$9::work_order_kind)
        AND ($10::uuid IS NULL OR w.id=$10::uuid)
        AND ($11::uuid IS NULL OR w.assignee_id=$11::uuid OR w.crew_id=$12::uuid)"""
    if rtype == "materials":
        rows = await conn.fetch(
            f"""SELECT m.code, m.name, sum(mw.quantity)::text AS quantity, mw.unit
                FROM material_writeoffs mw
                JOIN materials m ON m.id=mw.material_id
                JOIN work_orders w ON w.id=mw.work_order_id
                WHERE {where}
                GROUP BY m.code, m.name, mw.unit
                ORDER BY m.name, mw.unit LIMIT 10000""",
            *args,
        )
        return (
            ["code", "material", "quantity", "unit"],
            [[r["code"], r["name"], r["quantity"], r["unit"]] for r in rows],
        )
    if rtype == "downtime":
        rows = await conn.fetch(
            f"""SELECT w.id, w.number, e.name AS equipment, w.equipment_stopped_at,
                       w.equipment_restored_at
                FROM work_orders w JOIN equipment e ON e.id=w.equipment_id
                WHERE {where} AND w.equipment_stopped_at IS NOT NULL
                ORDER BY w.equipment_stopped_at DESC LIMIT 10000""",
            *args,
        )
        return (
            ["id", "number", "equipment", "stopped_at", "restored_at"],
            [
                [
                    str(r["id"]),
                    r["number"],
                    r["equipment"],
                    str(r["equipment_stopped_at"]),
                    str(r["equipment_restored_at"]),
                ]
                for r in rows
            ],
        )
    if rtype == "ratings":
        rows = await conn.fetch(
            f"""SELECT COALESCE(e.full_name, c.name) AS employee_or_crew,
                       count(*) FILTER (WHERE w.status='closed')::int AS closed,
                       count(*)::int AS total,
                       round(avg(COALESCE(ar.master_score, ar.score))::numeric, 1)::text AS average_review,
                       round(avg(CASE WHEN w.done_at IS NOT NULL AND w.done_at <= w.due_at
                                      THEN 100.0 ELSE 0.0 END)::numeric, 1)::text AS on_time_percent
                FROM work_orders w
                LEFT JOIN employees e ON e.id=w.assignee_id
                LEFT JOIN crews c ON c.id=w.crew_id
                LEFT JOIN LATERAL (
                    SELECT score, master_score FROM ai_reviews
                    WHERE work_order_id=w.id ORDER BY attempt DESC LIMIT 1
                ) ar ON true
                WHERE {where}
                GROUP BY e.id, e.full_name, c.id, c.name
                ORDER BY closed DESC, COALESCE(e.full_name, c.name) LIMIT 10000""",
            *args,
        )
        return (
            ["employee_or_crew", "closed", "total", "average_review", "on_time_percent"],
            [
                [
                    r["employee_or_crew"],
                    r["closed"],
                    r["total"],
                    r["average_review"],
                    r["on_time_percent"],
                ]
                for r in rows
            ],
        )
    if rtype == "dashboard":
        rows = await conn.fetch(
            f"""SELECT w.status::text AS status, count(*)::int AS orders,
                       count(*) FILTER (WHERE w.due_at < now()
                         AND w.status NOT IN ('done','ai_review','closed','rejected','cancelled'))::int AS overdue
                FROM work_orders w WHERE {where}
                GROUP BY w.status ORDER BY w.status LIMIT 10000""",
            *args,
        )
        return (
            ["status", "orders", "overdue"],
            [[r["status"], r["orders"], r["overdue"]] for r in rows],
        )
    rows = await conn.fetch(
        f"""SELECT w.id, w.number, w.status::text AS status, w.priority::text AS priority,
                   w.kind::text AS kind, a.name AS area, e.name AS equipment,
                   COALESCE(emp.full_name, c.name) AS assignee, w.description,
                   w.due_at, w.issued_at, w.closed_at
            FROM work_orders w
            JOIN areas a ON a.id=w.area_id
            JOIN equipment e ON e.id=w.equipment_id
            LEFT JOIN employees emp ON emp.id=w.assignee_id
            LEFT JOIN crews c ON c.id=w.crew_id
            WHERE {where}
            ORDER BY w.issued_at DESC, w.id LIMIT 10000""",
        *args,
    )
    headers = [
        "id",
        "number",
        "status",
        "priority",
        "kind",
        "area",
        "equipment",
        "assignee",
        "description",
        "due_at",
        "issued_at",
        "closed_at",
    ]
    return (headers, [[str(r[h]) if r[h] is not None else "" for h in headers] for r in rows])


def _xlsx(header: list[str], rows: list[list[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="194C93", fill_type="solid")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_font_candidates() -> list[str]:
    env = os.environ.get("PDF_FONT_PATH", "")
    if env:
        # Явно заданный путь — только он (ошибка конфигурации должна падать, а не молча подменяться).
        return [env]
    return [
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",  # Dockerfile: fonts-noto-core
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",  # generic linux fallback
        "/System/Library/Fonts/Supplemental/Arial.ttf",  # macOS local dev
    ]


def _pdf(header: list[str], rows: list[list[Any]]) -> bytes:
    candidates = _pdf_font_candidates()
    font_path = next((p for p in candidates if os.path.exists(p)), None)
    if font_path is None:
        raise RuntimeError(
            f"PDF_FONT_PATH must point to a Unicode TrueType font (tried: {', '.join(candidates)})"
        )
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("Noto", "", font_path)
    pdf.set_font("Noto", size=10)
    pdf.cell(0, 10, " | ".join(header))
    pdf.ln()
    pdf.set_font("Noto", size=8)
    for r in rows[:500]:
        pdf.cell(0, 6, " | ".join(str(x) for x in r)[:180])
        pdf.ln()
    return bytes(pdf.output())


async def process_export(conn: Any, export_id: str) -> None:
    from app.storage import put_bytes

    exp = await conn.fetchrow("SELECT * FROM report_exports WHERE id=$1::uuid", export_id)
    if not exp or exp["state"] == "completed":
        return
    await conn.execute("UPDATE report_exports SET state='processing' WHERE id=$1::uuid", export_id)
    try:
        header, rows = await _rows(conn, dict(exp))
        fmt = exp["format"]
        data = _xlsx(header, rows) if fmt == "xlsx" else _pdf(header, rows)
        key = f"reports/{exp['requested_by']}/{export_id}.{'xlsx' if fmt == 'xlsx' else 'pdf'}"
        put_bytes(
            key,
            data,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if fmt == "xlsx"
            else "application/pdf",
        )
        await conn.execute(
            "UPDATE report_exports SET state='completed', object_key=$2, completed_at=now() WHERE id=$1::uuid",
            export_id,
            key,
        )
    except Exception as e:
        await conn.execute(
            "UPDATE report_exports SET state='failed', error=$2 WHERE id=$1::uuid",
            export_id,
            str(e)[:500],
        )
