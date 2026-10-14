"""Report query and file-generation tests with controlled DB and object-store seams."""

import asyncio
import io
import sys
import uuid
from datetime import UTC, datetime
from types import SimpleNamespace

import pytest
from pydantic import ValidationError

from app.services import reports


def run(coro):
    return asyncio.run(coro)


class FakeConn:
    def __init__(self, row=None, rows=None):
        self.row = row
        self.rows = rows or []
        self.fetch_calls = []
        self.fetchrow_calls = []
        self.executions = []

    async def fetch(self, query, *args):
        self.fetch_calls.append((query, args))
        return self.rows

    async def fetchrow(self, query, *args):
        self.fetchrow_calls.append((query, args))
        return self.row

    async def execute(self, query, *args):
        self.executions.append((query, args))


def test_master_scope_and_all_supported_filters_are_bound_for_order_export():
    area_id, equipment_id, assignee_id, order_id = (uuid.uuid4() for _ in range(4))
    from_at = "2026-09-01T00:00:00Z"
    to_at = "2026-10-01T00:00:00Z"
    conn = FakeConn(
        rows=[
            {
                "id": order_id,
                "number": 12,
                "status": "closed",
                "priority": "high",
                "kind": "unplanned",
                "area": "A",
                "equipment": "Pump",
                "assignee": "Worker",
                "description": "Bearing",
                "due_at": None,
                "issued_at": None,
                "closed_at": None,
            }
        ]
    )

    headers, rows = run(
        reports._rows(
            conn,
            {
                "report_type": "orders",
                "filters": {
                    "from": from_at,
                    "to": to_at,
                    "area_id": str(area_id),
                    "equipment_id": str(equipment_id),
                    "assignee_id": str(assignee_id),
                    "work_order_id": str(order_id),
                    "status": "closed",
                    "priority": "high",
                    "kind": "unplanned",
                },
                "scope": {"role": "master", "area_ids": [str(area_id)]},
            },
        )
    )

    query, args = conn.fetch_calls[0]
    assert headers == [
        "id",
        "number",
        "status",
        "priority",
        "kind",
        "area",
        "equipment",
        "assignee",
        "description",
        "due_at",
        "issued_at",
        "closed_at",
    ]
    assert rows == [
        [
            str(order_id),
            "12",
            "closed",
            "high",
            "unplanned",
            "A",
            "Pump",
            "Worker",
            "Bearing",
            "",
            "",
            "",
        ]
    ]
    assert "w.area_id=ANY($1::uuid[])" in query
    assert "w.equipment_id=$5::uuid" in query
    assert "w.id=$10::uuid" in query
    assert "$11::uuid IS NULL" in query
    assert args == (
        [area_id],
        datetime(2026, 9, 1, tzinfo=UTC),
        datetime(2026, 10, 1, tzinfo=UTC),
        area_id,
        equipment_id,
        assignee_id,
        "closed",
        "high",
        "unplanned",
        order_id,
        None,
        None,
    )


def test_master_with_corrupt_or_empty_scope_fails_closed():
    conn = FakeConn()
    run(
        reports._rows(
            conn,
            {
                "report_type": "orders",
                "filters": {},
                "scope": {"role": "master", "area_ids": ["not-a-uuid"]},
            },
        )
    )
    query, args = conn.fetch_calls[-1]
    assert "w.area_id=ANY($1::uuid[])" in query
    assert args[0] == []

    run(
        reports._rows(
            conn,
            {"report_type": "orders", "filters": {}, "scope": {"role": "master", "area_ids": []}},
        )
    )
    assert conn.fetch_calls[-1][1][0] == []


def test_worker_scope_is_enforced_even_for_non_order_report_types():
    worker_id, crew_id = uuid.uuid4(), uuid.uuid4()
    conn = FakeConn(rows=[{"code": "F-1", "name": "Фильтр", "quantity": "2.5", "unit": "шт"}])

    headers, rows = run(
        reports._rows(
            conn,
            {
                "report_type": "materials",
                "filters": {},
                "scope": {"role": "worker", "user_id": str(worker_id), "crew_id": str(crew_id)},
            },
        )
    )

    query, args = conn.fetch_calls[0]
    assert headers == ["code", "material", "quantity", "unit"]
    assert rows == [["F-1", "Фильтр", "2.5", "шт"]]
    assert "JOIN work_orders w ON w.id=mw.work_order_id" in query
    assert "w.assignee_id=$11::uuid OR w.crew_id=$12::uuid" in query
    assert args[-2:] == (worker_id, crew_id)


def test_worker_missing_identity_fails_closed_and_admin_scope_stays_global():
    conn = FakeConn()
    run(
        reports._rows(
            conn,
            {"report_type": "orders", "filters": {}, "scope": {"role": "worker", "crew_id": None}},
        )
    )
    assert conn.fetch_calls[-1][1][-2] == uuid.UUID(int=0)

    run(
        reports._rows(
            conn,
            {
                "report_type": "orders",
                "filters": {},
                "scope": {"role": "admin", "user_id": str(uuid.uuid4())},
            },
        )
    )
    assert conn.fetch_calls[-1][1][0] is None
    assert conn.fetch_calls[-1][1][-2:] == (None, None)


def test_aggregate_reports_keep_work_order_scope_and_rating_crew_name():
    conn = FakeConn(
        rows=[
            {
                "employee_or_crew": "Бригада 7",
                "closed": 3,
                "total": 4,
                "average_review": "86.5",
                "on_time_percent": "75.0",
            }
        ]
    )
    headers, rows = run(
        reports._rows(
            conn,
            {
                "report_type": "ratings",
                "filters": {},
                "scope": {"role": "master", "area_ids": [str(uuid.uuid4())]},
            },
        )
    )
    query, args = conn.fetch_calls[0]
    assert headers[0] == "employee_or_crew"
    assert rows == [["Бригада 7", 3, 4, "86.5", "75.0"]]
    assert "COALESCE(e.full_name, c.name) AS employee_or_crew" in query
    assert "WHERE work_order_id=w.id ORDER BY attempt DESC LIMIT 1" in query
    assert args[0] and len(args) == 12

    conn.rows = [{"status": "issued", "orders": 2, "overdue": 1}]
    headers, rows = run(
        reports._rows(conn, {"report_type": "dashboard", "filters": {}, "scope": {"role": "admin"}})
    )
    assert headers == ["status", "orders", "overdue"]
    assert rows == [["issued", 2, 1]]


def test_material_and_downtime_exports_have_stable_columns_and_scope():
    conn = FakeConn(rows=[{"code": "M-1", "name": "Фильтр", "quantity": "2.5", "unit": "шт"}])
    headers, rows = run(
        reports._rows(conn, {"report_type": "materials", "filters": {}, "scope": {"role": "admin"}})
    )
    assert headers == ["code", "material", "quantity", "unit"]
    assert rows == [["M-1", "Фильтр", "2.5", "шт"]]
    assert "JOIN work_orders w ON w.id=mw.work_order_id" in conn.fetch_calls[-1][0]

    conn.rows = [
        {
            "id": uuid.uuid4(),
            "number": 7,
            "equipment": "Компрессор",
            "equipment_stopped_at": datetime(2026, 1, 1, tzinfo=UTC),
            "equipment_restored_at": None,
        }
    ]
    headers, rows = run(
        reports._rows(
            conn, {"report_type": "downtime", "filters": "{}", "scope": {"role": "admin"}}
        )
    )
    assert headers == ["id", "number", "equipment", "stopped_at", "restored_at"]
    assert rows[0][:3] == [str(conn.rows[0]["id"]), 7, "Компрессор"]
    assert rows[0][-1] == "None"
    assert "equipment_stopped_at IS NOT NULL" in conn.fetch_calls[-1][0]


def test_invalid_persisted_filters_are_rejected_before_query():
    conn = FakeConn()
    with pytest.raises(ValidationError):
        run(
            reports._rows(
                conn,
                {
                    "report_type": "orders",
                    "filters": {"status": "sql injection"},
                    "scope": {"role": "admin"},
                },
            )
        )
    with pytest.raises(ValidationError):
        run(
            reports._rows(
                conn,
                {
                    "report_type": "orders",
                    "filters": {"from": "2026-10-02T00:00:00Z", "to": "2026-10-01T00:00:00Z"},
                    "scope": {"role": "admin"},
                },
            )
        )
    with pytest.raises(ValueError, match="scope"):
        run(reports._rows(conn, {"report_type": "orders", "filters": {}, "scope": []}))
    with pytest.raises(ValueError, match="role"):
        run(
            reports._rows(
                conn, {"report_type": "orders", "filters": {}, "scope": {"role": "unknown"}}
            )
        )
    with pytest.raises(ValueError, match="role"):
        run(reports._rows(conn, {"report_type": "orders", "filters": {}}))
    assert conn.fetch_calls == []


def test_xlsx_is_a_readable_workbook_with_frozen_header_and_rows():
    from openpyxl import load_workbook

    content = reports._xlsx(["name", "qty"], [["Подшипник", 2]])
    wb = load_workbook(io.BytesIO(content), read_only=False)
    ws = wb.active
    assert list(ws.values) == [("name", "qty"), ("Подшипник", 2)]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.fgColor.rgb.endswith("194C93")


def test_pdf_requires_an_explicit_unicode_font(monkeypatch, tmp_path):
    monkeypatch.setenv("PDF_FONT_PATH", str(tmp_path / "absent.ttf"))
    with pytest.raises(RuntimeError, match="Unicode TrueType font"):
        reports._pdf(["id"], [["1"]])


def test_pdf_uses_unicode_font_limits_rows_and_truncates_oversized_cells(monkeypatch):
    class FakePDF:
        def __init__(self):
            self.cells = []
            self.fonts = []
            self.font_choices = []

        def add_page(self):
            self.page_added = True

        def add_font(self, *args):
            self.fonts.append(args)

        def set_font(self, *args, **kwargs):
            self.font_choices.append((args, kwargs))

        def cell(self, width, height, text):
            self.cells.append((width, height, text))

        def ln(self):
            pass

        def output(self):
            return bytearray(b"%PDF-test")

    fake_pdf = FakePDF()
    monkeypatch.setenv("PDF_FONT_PATH", "/fonts/NotoSans-Regular.ttf")
    monkeypatch.setattr(reports.os.path, "exists", lambda _path: True)
    monkeypatch.setitem(sys.modules, "fpdf", SimpleNamespace(FPDF=lambda: fake_pdf))

    rows = [[str(index), "Ж" * 250] for index in range(501)]
    content = reports._pdf(["номер", "описание"], rows)

    assert content == b"%PDF-test"
    assert fake_pdf.page_added
    assert fake_pdf.fonts == [("Noto", "", "/fonts/NotoSans-Regular.ttf")]
    assert fake_pdf.font_choices == [(("Noto",), {"size": 10}), (("Noto",), {"size": 8})]
    assert fake_pdf.cells[0][2] == "номер | описание"
    row_cells = fake_pdf.cells[1:]
    assert len(row_cells) == 500
    assert row_cells[0][2] == "0 | " + "Ж" * 176
    assert row_cells[-1][2].startswith("499 | ")


def test_process_export_is_idempotent_for_missing_and_completed_rows(monkeypatch):
    conn = FakeConn()
    run(reports.process_export(conn, "missing"))
    assert conn.executions == []

    conn.row = {"state": "completed"}
    run(reports.process_export(conn, "done"))
    assert conn.executions == []
    assert len(conn.fetchrow_calls) == 2


def test_process_export_uploads_xlsx_then_marks_completed(monkeypatch):
    conn = FakeConn(
        row={
            "state": "queued",
            "report_type": "shift",
            "filters": {},
            "scope": {},
            "format": "xlsx",
            "requested_by": "user-1",
        }
    )
    writes = []
    monkeypatch.setattr(reports, "_rows", lambda *_: asyncio.sleep(0, result=(["id"], [["1"]])))
    monkeypatch.setattr("app.storage.put_bytes", lambda *args: writes.append(args))

    run(reports.process_export(conn, "export-1"))

    assert conn.executions[0][0].startswith("UPDATE report_exports SET state='processing'")
    assert conn.executions[-1][0].startswith("UPDATE report_exports SET state='completed'")
    assert conn.executions[-1][1] == ("export-1", "reports/user-1/export-1.xlsx")
    assert writes[0][0] == "reports/user-1/export-1.xlsx"
    assert writes[0][2] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert writes[0][1][:2] == b"PK"


def test_process_export_persists_generation_failure_for_operator_visibility(monkeypatch):
    conn = FakeConn(
        row={
            "state": "queued",
            "report_type": "shift",
            "filters": {},
            "scope": {},
            "format": "pdf",
            "requested_by": "user-1",
        }
    )
    monkeypatch.setattr(reports, "_rows", lambda *_: asyncio.sleep(0, result=(["id"], [])))
    monkeypatch.setattr(
        reports, "_pdf", lambda *_: (_ for _ in ()).throw(RuntimeError("font missing"))
    )

    run(reports.process_export(conn, "export-2"))

    assert conn.executions[-1][0].startswith("UPDATE report_exports SET state='failed'")
    assert conn.executions[-1][1] == ("export-2", "font missing")
